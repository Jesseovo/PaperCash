"""Offline regression tests for the CLI's substantive behavior."""

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "scripts" / "lib"))

import papercash
from features.ai_humanize import rewrite_text
from features.lit_review import generate_review
from features.docx_export import export_review_to_docx, export_search_to_docx
from features import citation
from features import plagiarism_v2
from dedupe import dedupe_papers
from schema import Paper, SearchResult


def test_rewrite_preserves_sentence_with_dose_and_citation():
    original = "综上所述，患者服用阿司匹林100mg，结果见文献[1]。综上所述，这一方法值得讨论。"
    modified = rewrite_text(original)
    assert modified.startswith("综上所述，患者服用阿司匹林100mg，结果见文献[1]。")
    assert "综合上述分析，这一方法值得讨论。" in modified
    assert rewrite_text("根据《民法典》第123条规定，综上所述。") == "根据《民法典》第123条规定，综上所述。"


def test_check_without_candidates_is_unverified(monkeypatch):
    monkeypatch.setattr(plagiarism_v2, "_search_multiple_sources", lambda *args, **kwargs: [])
    results, stats = plagiarism_v2.check_plagiarism_v2("本研究采用可重复的实验方法分析样本差异。", mode="fast")
    assert len(results) == 1
    assert results[0].risk_level == "unverified"
    assert stats["unverified"] == 1
    assert stats["compared_sentences"] == 0
    assert "accuracy_estimate" not in stats


def test_check_compares_against_abstract_sentence(monkeypatch):
    sentence = "本研究采用可重复的实验方法分析样本差异"
    paper = Paper(title="Unrelated title", abstract=f"This is a long introduction. {sentence}。后续内容。")
    monkeypatch.setattr(plagiarism_v2, "_search_multiple_sources", lambda *args, **kwargs: [paper])
    results, stats = plagiarism_v2.check_plagiarism_v2(sentence + "。", mode="fast")
    assert stats["compared_sentences"] == 1
    assert results[0].matched_paper is paper
    assert results[0].similarity["literal"] > 0.9
    assert results[0].similarity["detection_type"] == "near_exact"


def test_review_labels_metadata_and_single_reference_number():
    paper = Paper(title="Example study", authors=["Jane Doe"], year=2024,
                  abstract="An example abstract.", doi="10.1234/example")
    review = generate_review(SearchResult(query="example", papers=[paper],
                                           sources_used=["fixture"], total_found=1))
    assert "仅依据检索元数据和可用摘要" in review
    assert "研究空白与未来方向" not in review
    assert "[1] [Doe2024]" not in review


def test_humanize_json_is_parseable_and_protected(monkeypatch, capsys):
    import safety

    monkeypatch.setattr(safety, "require_explicit_confirmation",
                        lambda feature: {"confirmed": True, "purpose": "测试"})
    monkeypatch.setattr(safety, "log_operation", lambda *args: None)
    monkeypatch.setattr(safety, "check_abuse", lambda: None)
    papercash.cmd_humanize(argparse.Namespace(
        input="综上所述，剂量为100mg。综上所述，该问题值得讨论。",
        emit="json", allow_modify_protected=False))
    output = json.loads(capsys.readouterr().out)
    assert output["protected"] is True
    assert "综上所述，剂量为100mg。" in output["suggestion"]
    assert "综合上述分析，该问题值得讨论。" in output["suggestion"]


def test_confirmation_requires_matching_code(monkeypatch):
    import safety

    monkeypatch.setattr(safety.random, "randint", lambda low, high: 1234)
    answers = iter(["0000"])
    monkeypatch.setattr("builtins.input", lambda: next(answers))
    assert safety.require_explicit_confirmation("humanize")["confirmed"] is False

    answers = iter(["1234", "1"])
    monkeypatch.setattr("builtins.input", lambda: next(answers))
    result = safety.require_explicit_confirmation("humanize")
    assert result["confirmed"] is True
    assert result["purpose"] == "学习"


def test_search_json_is_parseable(monkeypatch, capsys):
    result = SearchResult(query="example", papers=[], total_found=0, sources_used=[])
    monkeypatch.setattr(papercash, "_search_all_sources", lambda *args, **kwargs: result)
    papercash.cmd_search(argparse.Namespace(query="example", limit=2, sources=None,
                                         emit="json", save_dir=None, export_docx=None))
    assert json.loads(capsys.readouterr().out)["query"] == "example"


def test_dedupe_keeps_doi_index_after_title_match():
    first = Paper(title="A detailed study of graphs", doi="10.1/first", final_score=90)
    second = Paper(title="A detailed study of graphs", doi="10.1/second", final_score=80,
                   abstract="Additional metadata")
    third = Paper(title="A detailed study of graphs", doi="10.1/second", final_score=70)
    result = dedupe_papers([first, second, third])
    assert len(result) == 1
    assert result[0].abstract == "Additional metadata"


def test_dedupe_does_not_merge_distinct_titles_with_same_prefix():
    prefix = "A study of graph neural networks for medical imaging"
    papers = [Paper(title=prefix + " in hospitals"),
              Paper(title=prefix + " in clinics")]
    assert len(dedupe_papers(papers)) == 2


def test_word_exports_open_with_expected_content(tmp_path):
    from docx import Document

    paper = Paper(title="A verifiable paper", authors=["Jane Doe"], doi="10.1234/example")
    result = SearchResult(query="example", papers=[paper], total_found=1,
                          sources_used=["fixture"])
    search_path = tmp_path / "search.docx"
    review_path = tmp_path / "review.docx"
    assert export_search_to_docx(result, str(search_path)).startswith("已导出")
    assert export_review_to_docx(generate_review(result), str(review_path)).startswith("已导出")
    assert "10.1234/example" in "\n".join(p.text for p in Document(search_path).paragraphs)
    assert "后续核查" in "\n".join(p.text for p in Document(review_path).paragraphs)


def test_citation_uses_available_publication_metadata(monkeypatch):
    paper = Paper(title="Example study", authors=["Jane Doe"], year=2024,
                  venue="Journal of Examples", volume="12", issue="2", pages="10-20",
                  doi="10.1234/example")
    monkeypatch.setattr(citation, "_get_paper_from_doi", lambda doi: paper)
    gb = citation.format_citation("10.1234/example", "gb7714")
    apa = citation.format_citation("10.1234/example", "apa")
    assert gb.startswith("[1] Jane Doe.")
    assert "2024, 12(2): 10-20" in gb
    assert "*Journal of Examples, 12*(2), 10-20" in apa
    assert citation.format_citations_batch(["one", "two"], "gb7714").count("[1]") == 1
